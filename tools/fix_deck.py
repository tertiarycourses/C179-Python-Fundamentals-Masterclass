"""Post-conversion fix-ups for the C179 deck: strip the WSQ layer convert_deck.py left behind."""
import re
from pptx import Presentation

PATH = "courseware/Python Fundamentals Masterclass-v1.0.pptx"
p = Presentation(PATH)


def find(i, startswith):
    for sh in p.slides[i].shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(startswith):
            return sh
    raise SystemExit(f"slide {i}: no shape starting {startswith!r}")


def set_paras(sh, lines):
    """Write one line per existing paragraph (through runs[0]); drop surplus paragraphs."""
    paras = sh.text_frame.paragraphs
    assert len(paras) >= len(lines), (sh.shape_id, len(paras), len(lines))
    for para, line in zip(paras, lines):
        para.runs[0].text = line
        for r in para.runs[1:]:
            r.text = ""
    for para in paras[len(lines):]:
        para._p.getparent().remove(para._p)


def set_para(sh, k, line):
    para = sh.text_frame.paragraphs[k]
    para.runs[0].text = line
    for r in para.runs[1:]:
        r.text = ""


# Cover
set_paras(find(0, "COURSE SLIDES"), ["COURSE SLIDES  ·  NON-WSQ COURSE"])
set_paras(find(0, "WSQ Course Code"), [
    "Course Code: C179",
    "Duration: 2 days  ·  15 instructional hours (9:30am – 5:30pm)",
    "Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W"])

# Trainer profile
set_para(find(3, "DELIVERS"), 1, "Courses on Python programming, data science and machine learning")

# LMS slide: open-your-course step carries the title already (convert_deck); nothing WSQ left.

# Schedule
set_paras(find(7, "•  Day 1"), [
    "•  Day 1 — Python Basics — Data Types, Operators, Control Structures",
    "–  Welcome & introductions",
    "–  Topic 1: Get Started on Python",
    "–  Topic 2: Data Types",
    "–  Lunch (30 min)",
    "–  Topic 3: Operators",
    "–  Topic 4: Control Structures, Loop & Comprehension"])
set_paras(find(7, "•  Day 2"), [
    "•  Day 2 — Functions, Modules & Packages",
    "–  Day 1 recap",
    "–  Topic 5: Function",
    "–  Lunch (30 min)",
    "–  Topic 6: Modules & Packages",
    "–  Course feedback, wrap-up and Q&A",
    "–  Daily timing: 9:30am – 5:30pm · 30-min lunch"])

# Course outline: drop the "(K1, A1)" competency chips
for i, s in enumerate(p.slides):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            new = re.sub(r"\s*\((?:[KA]\d+(?:[–-][KA]?\d+)?(?:,\s*)?)+\)$", "", t)
            if new != t:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""

# Eyebrows "TSC K1, A1" -> "KEY CONCEPTS"
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame and len(sh.text_frame.paragraphs) == 1 and re.match(r"^(TSC|C179) [KA]\d", sh.text_frame.text):
            set_paras(sh, ["KEY CONCEPTS"])

# Recommended next courses -> live non-WSQ courses
rec = next(i for i, s in enumerate(p.slides)
           if any(sh.has_text_frame and sh.text_frame.text == "Recommended Next Courses" for sh in s.shapes))
for start, name in [("Data Analytics and Deep", "Python Data Analysis Masterclass"),
                    ("R Fundamental", "AI Vibe Coding with Python"),
                    ("Advancing Your Python", "SQL Masterclass"),
                    ("Building Your First Machine", "AI Vibe Coding for Machine Learning"),
                    ("Search Engine Optimization", "AI Vibe Coding for Deep Learning")]:
    sh = find(rec, start)
    set_para(sh, 0, name)
    set_para(sh, 1, "Hands-on course · Tertiary Courses")

# Course Feedback (rewritten TRAQOM slide)
fb = next(i for i, s in enumerate(p.slides)
          if any(sh.has_text_frame and sh.text_frame.text == "FEEDBACK & CERTIFICATE" for sh in s.shapes))
set_paras(find(fb, "Cert &") if any(sh.has_text_frame and sh.text_frame.text.startswith("Cert &") for sh in p.slides[fb].shapes)
          else find(fb, "Course Feedback"), ["Course Feedback & Certificate"])
set_paras(find(fb, "•  "), [
    "•  Share your feedback in the course survey on the LMS — it helps us improve the course.",
    "•  Your certificate of completion is issued via the LMS after the course.",
    "•  Portal: https://lms-tms.tertiaryinfotech.com/"])

p.save(PATH)
print("saved; feedback slide", fb, "recommended slide", rec)

# --- second pass: the Day 2 card lost its pPr when convert_deck's --sub rewrote it;
#     copy paragraph properties from the untouched Day 1 card, paragraph by paragraph.
import copy
p = Presentation(PATH)
d1, d2 = find(7, "•  Day 1"), find(7, "•  Day 2")
for k, para in enumerate(d2.text_frame.paragraphs):
    src = d1.text_frame.paragraphs[min(k, len(d1.text_frame.paragraphs) - 1)]._p
    src_ppr, src_rpr = src.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pPr'), src.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
    old = para._p.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pPr')
    if old is not None:
        para._p.remove(old)
    if src_ppr is not None:
        para._p.insert(0, copy.deepcopy(src_ppr))
    for r in para.runs:
        rp = r._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        if rp is not None:
            r._r.remove(rp)
        if src_rpr is not None:
            r._r.insert(0, copy.deepcopy(src_rpr))
p.save(PATH)
print("day-2 card formatting restored")
